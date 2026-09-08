import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Excel dosyasını oku
df = pd.read_excel("tehlike_listesi.xlsx")
df.columns = df.columns.str.strip()

# 2. CCP Yönetim Parametreleri
ccp_kurallari = {
    "Haslama": {
        "Kritik_Limit": "Sicaklik >= 90 C, Sure >= 90 sn",
        "Izleme_Yontemi": "Hat ici PT100 sensoru ile surekli sicaklik kaydi",
        "Izleme_Sikligi": "Surekli (Otomatik SCADA)",
        "Duzeltici_Faaliyet": "Limit alti urun tahliye edilir, karantinaya alinir, yeniden haslanir."
    },
    "Paketleme": {
        "Kritik_Limit": "Fe: 1.5 mm, Non-Fe: 2.0 mm, SS: 2.5 mm",
        "Izleme_Yontemi": "Test cubuklari ile dedektor sinyal testi",
        "Izleme_Sikligi": "Vardiya basi, sonu ve saatte 1",
        "Duzeltici_Faaliyet": "Dedektor durdurulur, son 1 saatlik uretim bloke edilir ve yeniden elenir."
    }
}

# 3. Analiz ve Kural Entegrasyonu
def haccp_yonetimi(row):
    proses = str(row["Proses_Adimi"]).strip()
    s1 = str(row["S1_Onlem_Var_Mi"]).strip().upper() == "EVET"
    s2 = str(row["S2_Yok_Ediyor_Mu"]).strip().upper() == "EVET"
    s3 = str(row["S3_Artis_Var_Mi"]).strip().upper() == "EVET"
    s4 = str(row["S4_Sonra_Yok_Olur_Mu"]).strip().upper() == "EVET"

    if not s1:
        karar = "Tasarim Revizyonu"
    elif s2 or (s3 and not s4):
        karar = "CCP"
    elif s4:
        karar = "oPRP"
    else:
        karar = "PRP"

    if karar == "CCP" and proses in ccp_kurallari:
        limit = ccp_kurallari[proses]["Kritik_Limit"]
        izleme = ccp_kurallari[proses]["Izleme_Yontemi"]
        siklik = ccp_kurallari[proses]["Izleme_Sikligi"]
        faaliyet = ccp_kurallari[proses]["Duzeltici_Faaliyet"]
    else:
        limit = "-"
        izleme = "Standart Hijyen / Proses Kontrolu"
        siklik = "Gunluk / Rutin Kontrol"
        faaliyet = "Standart Prosedur"

    return pd.Series([karar, limit, izleme, siklik, faaliyet])

df[["HACCP_Karari", "Kritik_Limit", "Izleme_Yontemi", "Izleme_Sikligi", "Duzeltici_Faaliyet"]] = df.apply(haccp_yonetimi, axis=1)

# 4. Veriyi Kaydet
cikti_dosyasi = "haccp_plani_final.xlsx"
df.to_excel(cikti_dosyasi, index=False)

# 5. openpyxl ile Profesyonel Stil ve Biçimlendirme
wb = load_workbook(cikti_dosyasi)
ws = wb.active
ws.title = "HACCP Plani"

# Renk tanımları
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Koyu Mavi/Lacivert
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

fill_ccp = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")   # Açık Kırmızı / Somon
font_ccp = Font(name="Calibri", size=10, bold=True, color="C00000")

fill_oprp = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # Açık Sarı
font_oprp = Font(name="Calibri", size=10, bold=True, color="B25900")

fill_prp = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")   # Açık Yeşil
font_prp = Font(name="Calibri", size=10, bold=True, color="375623")

thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)

# Başlık stilini uygula
for col_num in range(1, ws.max_column + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws.row_dimensions[1].height = 28

# Satırları ve karara göre renklendirmeyi uygula
for row_num in range(2, ws.max_row + 1):
    ws.row_dimensions[row_num].height = 24
    karar_cell = ws.cell(row=row_num, column=df.columns.get_loc("HACCP_Karari") + 1)
    karar_degeri = str(karar_cell.value).strip()

    secilen_fill = None
    secilen_font = None

    if "CCP" in karar_degeri:
        secilen_fill = fill_ccp
        secilen_font = font_ccp
    elif "oPRP" in karar_degeri:
        secilen_fill = fill_oprp
        secilen_font = font_oprp
    elif "PRP" in karar_degeri:
        secilen_fill = fill_prp
        secilen_font = font_prp

    for col_num in range(1, ws.max_column + 1):
        cell = ws.cell(row=row_num, column=col_num)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        
        # Karar hücresine özel vurgu
        if col_num == df.columns.get_loc("HACCP_Karari") + 1 and secilen_fill:
            cell.fill = secilen_fill
            cell.font = secilen_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

# Sütun genişliklerini otomatik ayarla
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.value:
            max_len = max(max_len, len(str(cell.value)))
    ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

wb.save(cikti_dosyasi)

print("\n--- PROFESYONEL HACCP RAPORU HAZIRLANDI ---")
print("Dosya: haccp_plani_final.xlsx")